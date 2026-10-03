from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.models import User
from django.contrib.auth import (
    authenticate,
    login as auth_login,
    logout,
    update_session_auth_hash
)
from django.contrib import messages
from django.contrib.auth.forms import PasswordChangeForm
from .models import Product, Sale, Customer, Supplier, StockMovement
from django.db.models import Sum, Count, Max, Avg
from django.utils import timezone
from openpyxl import Workbook
from django.http import HttpResponse
from django.core.paginator import Paginator
from django.contrib import messages
from django.core.exceptions import ValidationError


def dashboard(request):
    if not request.user.is_authenticated:
        return redirect('home')

    total_products = Product.objects.filter(
        user=request.user
    ).count()

    total_customers = Customer.objects.filter(
        user=request.user
    ).count()

    total_sales = Sale.objects.filter(
        user=request.user
    ).count()

    total_stock = Product.objects.filter(
        user=request.user
    ).aggregate(
        total=Sum('quantity')
    )['total'] or 0

    total_revenue = Sale.objects.filter(
        user=request.user
    ).aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    today = timezone.localdate()

    today_sales = Sale.objects.filter(
        user=request.user,
        sale_date__date=today
    ).aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    low_stock = Product.objects.filter(
        user=request.user,
        quantity__lte=5
    ).count()

    pending_payments = Sale.objects.filter(
        user=request.user,
        payment_status='Pending'
    ).count()

    recent_sales = Sale.objects.filter(
    user=request.user
    ).select_related(
    'product',
    'customer'
    ).order_by('-sale_date')[:5]

    low_stock_products = Product.objects.filter(
    user=request.user,
    quantity__lte=5
    ).order_by('quantity')[:5]

    last_7_days = []

    for i in range(6, -1, -1):
        day = timezone.localdate() - timezone.timedelta(days=i)

        amount = Sale.objects.filter(
            user=request.user,
            sale_date__date=day
        ).aggregate(
            total=Sum('total_amount')
        )['total'] or 0

        last_7_days.append({
            'date': day.strftime('%d %b'),
            'amount': float(amount)
        })

    return render(request, 'dashboard.html', {
        'total_products': total_products,
        'total_customers': total_customers,
        'total_sales': total_sales,
        'total_stock': total_stock,
        'total_revenue': total_revenue,
        'today_sales': today_sales,
        'low_stock': low_stock,
        'pending_payments': pending_payments,
        'recent_sales': recent_sales,
        'low_stock_products': low_stock_products,
        'last_7_days': last_7_days,
    })


def home(request):
    return render(request, 'home.html')


def login(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:
            auth_login(request, user)
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid username or password.')

    return render(request, 'login.html')


def register(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        if password != confirm_password:
            messages.error(request, 'Passwords do not match.')
            return render(request, 'register.html')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username already exists.')
            return render(request, 'register.html')

        user = User.objects.create_user(
            username=username,
            email=email,
            password=password
        )

        user.save()

        messages.success(request, 'Account created successfully. Please login.')
        return redirect('login')

    return render(request, 'register.html')


def logout_view(request):
    logout(request)
    return redirect('login')

def inventory(request):
    if not request.user.is_authenticated:
        return redirect('home')

    products = Product.objects.filter(
        user=request.user
    )

    search = request.GET.get('search')
    category = request.GET.get('category')

    if search:
        products = products.filter(
            name__icontains=search
        )

    if category:
        products = products.filter(
            category=category
        )

    categories = Product.objects.filter(
        user=request.user
    ).values_list(
        'category',
        flat=True
    ).distinct()

    # Pagination
    paginator = Paginator(products, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'inventory.html', {
        'products': page_obj,
        'categories': categories,
        'page_obj': page_obj
    })

def add_product(request):
    if not request.user.is_authenticated:
        return redirect('home')

    suppliers = Supplier.objects.filter(
        user=request.user
    )

    if request.method == 'POST':

        name = request.POST.get('name', '').strip()
        sku = request.POST.get('sku', '').strip()
        category = request.POST.get('category', '').strip()
        supplier_id = request.POST.get('supplier')
        quantity = request.POST.get('quantity', '').strip()
        purchase_price = request.POST.get('purchase_price', '').strip()
        selling_price = request.POST.get('selling_price', '').strip()

        errors = []

        # Required fields
        if not name:
            errors.append('Product name is required.')

        if not sku:
            errors.append('SKU is required.')

        if not category:
            errors.append('Category is required.')

        if not quantity:
            errors.append('Quantity is required.')

        if not purchase_price:
            errors.append('Purchase price is required.')

        if not selling_price:
            errors.append('Selling price is required.')

        # Number validation
        try:
            quantity = int(quantity)

            if quantity < 0:
                errors.append('Quantity cannot be negative.')

        except ValueError:
            errors.append('Quantity must be a valid number.')

        try:
            purchase_price = float(purchase_price)

            if purchase_price < 0:
                errors.append('Purchase price cannot be negative.')

        except ValueError:
            errors.append('Purchase price must be a valid number.')

        try:
            selling_price = float(selling_price)

            if selling_price < 0:
                errors.append('Selling price cannot be negative.')

        except ValueError:
            errors.append('Selling price must be a valid number.')

        # Check duplicate SKU
        if sku and Product.objects.filter(
            user=request.user,
            sku=sku
        ).exists():
            errors.append('A product with this SKU already exists.')

        # Supplier validation
        supplier = None

        if supplier_id:
            try:
                supplier = Supplier.objects.get(
                    id=supplier_id,
                    user=request.user
                )
            except Supplier.DoesNotExist:
                errors.append('Invalid supplier selected.')

        # Show errors
        if errors:
            for error in errors:
                messages.error(request, error)

            return render(request, 'add_product.html', {
                'suppliers': suppliers
            })

        # Create product
        product = Product.objects.create(
            user=request.user,
            name=name,
            sku=sku,
            category=category,
            supplier=supplier,
            quantity=quantity,
            purchase_price=purchase_price,
            selling_price=selling_price
        )

        # Record stock movement
        StockMovement.objects.create(
            user=request.user,
            product=product,
            movement_type='IN',
            quantity=product.quantity,
            previous_stock=0,
            new_stock=product.quantity,
            note='Initial stock added'
        )

        messages.success(
            request,
            'Product added successfully.'
        )

        return redirect('inventory')

    return render(request, 'add_product.html', {
        'suppliers': suppliers
    })


def edit_product(request, product_id):
    if not request.user.is_authenticated:
        return redirect('home')

    product = get_object_or_404(
        Product,
        id=product_id,
        user=request.user
    )

    supplier_list = Supplier.objects.filter(
        user=request.user
    ).order_by('name')

    if request.method == 'POST':
        product.name = request.POST.get('name')
        product.sku = request.POST.get('sku')
        product.category = request.POST.get('category')

        supplier_id = request.POST.get('supplier')

        if supplier_id:
            product.supplier = get_object_or_404(
                Supplier,
                id=supplier_id,
                user=request.user
            )
        else:
            product.supplier = None

        # Store old quantity
        old_quantity = product.quantity

        # Get new quantity
        new_quantity = int(request.POST.get('quantity'))

        product.quantity = new_quantity
        product.purchase_price = request.POST.get('purchase_price')
        product.selling_price = request.POST.get('selling_price')

        product.save()

        # Record stock movement if quantity changed
        if old_quantity != new_quantity:
            StockMovement.objects.create(
                user=request.user,
                product=product,
                movement_type='ADJUSTMENT',
                quantity=abs(new_quantity - old_quantity),
                previous_stock=old_quantity,
                new_stock=new_quantity,
                note='Stock adjusted while editing product'
            )

        messages.success(
            request,
            'Product updated successfully.'
        )

        return redirect('inventory')

    return render(request, 'edit_product.html', {
        'product': product,
        'suppliers': supplier_list
    })


def delete_product(request, product_id):
    if not request.user.is_authenticated:
        return redirect('home')

    product = get_object_or_404(
        Product,
        id=product_id,
        user=request.user
    )

    if request.method == 'POST':
        product.delete()
        messages.success(
            request, 'product deleted successfully.'
        )

        return redirect('inventory')

    return render(request, 'delete_product.html', {
        'product': product
    })

def sales(request):
    if not request.user.is_authenticated:
        return redirect('home')

    sales_list = Sale.objects.filter(
        user=request.user
    ).order_by('-sale_date')

    # Pagination
    paginator = Paginator(sales_list, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'sales.html', {
        'sales': page_obj,
        'page_obj': page_obj
    })

def new_sale(request):
    if not request.user.is_authenticated:
        return redirect('home')

    products = Product.objects.filter(
        user=request.user
    )

    customers = Customer.objects.filter(
        user=request.user
    )

    if request.method == 'POST':

        product_id = request.POST.get('product')
        customer_id = request.POST.get('customer')
        quantity_value = request.POST.get('quantity', '').strip()
        payment_status = request.POST.get('payment_status')

        # Required fields
        if not product_id:
            messages.error(request, 'Please select a product.')

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        if not customer_id:
            messages.error(request, 'Please select a customer.')

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        if not quantity_value:
            messages.error(request, 'Please enter a quantity.')

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        # Quantity validation
        try:
            quantity = int(quantity_value)

            if quantity <= 0:
                messages.error(
                    request,
                    'Quantity must be greater than 0.'
                )

                return render(request, 'new_sale.html', {
                    'products': products,
                    'customers': customers
                })

        except ValueError:
            messages.error(
                request,
                'Quantity must be a valid number.'
            )

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        # Product validation
        product = get_object_or_404(
            Product,
            id=product_id,
            user=request.user
        )

        # Stock validation
        if quantity > product.quantity:
            messages.error(
                request,
                f'Not enough stock available. '
                f'Only {product.quantity} item(s) available.'
            )

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        # Customer validation
        customer = get_object_or_404(
            Customer,
            id=customer_id,
            user=request.user
        )

        # Payment status validation
        allowed_statuses = [
            'Paid',
            'Pending',
            'Partial'
        ]

        if payment_status not in allowed_statuses:
            messages.error(
                request,
                'Invalid payment status.'
            )

            return render(request, 'new_sale.html', {
                'products': products,
                'customers': customers
            })

        selling_price = product.selling_price
        total_amount = selling_price * quantity

        # Store previous stock
        previous_stock = product.quantity

        # Create sale
        Sale.objects.create(
            user=request.user,
            product=product,
            customer=customer,
            customer_name=customer.name,
            quantity=quantity,
            selling_price=selling_price,
            total_amount=total_amount,
            payment_status=payment_status
        )

        # Update stock
        product.quantity -= quantity
        product.save()

        # Record stock movement
        StockMovement.objects.create(
            user=request.user,
            product=product,
            movement_type='OUT',
            quantity=quantity,
            previous_stock=previous_stock,
            new_stock=product.quantity,
            note='Stock reduced due to sale'
        )

        messages.success(
            request,
            'Sale created successfully.'
        )

        return redirect('sales')

    return render(request, 'new_sale.html', {
        'products': products,
        'customers': customers
    })


def invoices(request):
    if not request.user.is_authenticated:
        return redirect('home')

    invoice_list = Sale.objects.filter(
        user=request.user
    ).order_by('-sale_date')

    return render(request, 'invoices.html', {
        'invoices': invoice_list
    })

def invoice_detail(request, sale_id):
    if not request.user.is_authenticated:
        return redirect('home')

    sale = get_object_or_404(
        Sale,
        id=sale_id,
        user=request.user
    )

    return render(request, 'invoice_detail.html', {
        'sale': sale
    })

def customers(request):
    if not request.user.is_authenticated:
        return redirect('home')

    search = request.GET.get('search', '')

    customer_list = Customer.objects.filter(
        user=request.user
    )

    if search:
        customer_list = customer_list.filter(
            name__icontains=search
        )

    customer_list = customer_list.annotate(
        total_purchases=Sum('sale__total_amount')
    ).order_by('-created_at')

    # Pagination
    paginator = Paginator(customer_list, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'customers.html', {
        'customers': page_obj,
        'page_obj': page_obj
    })

def add_customer(request):
    if not request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':

        name = request.POST.get('name', '').strip()
        email = request.POST.get('email', '').strip()
        phone = request.POST.get('phone', '').strip()

        errors = []

        if not name:
            errors.append('Customer name is required.')

        if not phone:
            errors.append('Phone number is required.')

        if email and '@' not in email:
            errors.append('Please enter a valid email address.')

        if errors:
            for error in errors:
                messages.error(request, error)

            return render(request, 'add_customer.html')

        Customer.objects.create(
            user=request.user,
            name=name,
            email=email,
            phone=phone
        )

        messages.success(
            request,
            'Customer added successfully.'
        )

        return redirect('customers')

    return render(request, 'add_customer.html')



def edit_customer(request, customer_id):
    if not request.user.is_authenticated:
        return redirect('home')

    customer = get_object_or_404(
        Customer,
        id=customer_id,
        user=request.user
    )

    if request.method == 'POST':
        customer.name = request.POST.get('name')
        customer.email = request.POST.get('email')
        customer.phone = request.POST.get('phone')
        customer.address = request.POST.get('address')

        customer.save()

        messages.success(
            request,
            'Customer updated successfully.'
        )

        return redirect('customers')

    return render(request, 'edit_customer.html', {
        'customer': customer
    })

def delete_customer(request, customer_id):
    if not request.user.is_authenticated:
        return redirect('home')

    customer = get_object_or_404(
        Customer,
        id=customer_id,
        user=request.user
    )

    if request.method == 'POST':
        customer.delete()

        messages.success(
            request,
            'Customer deleted successfully.'
        )

        return redirect('customers')

    return render(request, 'delete_customer.html', {
        'customer': customer
    })


def suppliers(request):
    if not request.user.is_authenticated:
        return redirect('home')

    supplier_list = Supplier.objects.filter(
        user=request.user
    ).order_by('-created_at')

    # Pagination
    paginator = Paginator(supplier_list, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'suppliers.html', {
        'suppliers': page_obj,
        'page_obj': page_obj
    })

def add_supplier(request):
    if not request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':

        name = request.POST.get('name', '').strip()
        email = request.POST.get('email', '').strip()
        phone = request.POST.get('phone', '').strip()
        address = request.POST.get('address', '').strip()
        company = request.POST.get('company', '').strip()

        errors = []

        if not name:
            errors.append('Supplier name is required.')

        if not phone:
            errors.append('Phone number is required.')

        if email and '@' not in email:
            errors.append('Please enter a valid email address.')

        if errors:
            for error in errors:
                messages.error(request, error)

            return render(request, 'add_supplier.html')

        Supplier.objects.create(
            user=request.user,
            name=name,
            email=email,
            phone=phone,
            address=address,
            company=company
        )

        messages.success(
            request,
            'Supplier added successfully.'
        )

        return redirect('suppliers')

    return render(request, 'add_supplier.html')


def edit_supplier(request, supplier_id):
    if not request.user.is_authenticated:
        return redirect('home')

    supplier = get_object_or_404(
        Supplier,
        id=supplier_id,
        user=request.user
    )

    if request.method == 'POST':
        supplier.name = request.POST.get('name')
        supplier.company = request.POST.get('company')
        supplier.email = request.POST.get('email')
        supplier.phone = request.POST.get('phone')
        supplier.address = request.POST.get('address')

        supplier.save()

        messages.success(
            request,
            'Supplier updated successfully.'
        )

        return redirect('suppliers')

    return render(request, 'edit_supplier.html', {
        'supplier': supplier
    })


def delete_supplier(request, supplier_id):
    if not request.user.is_authenticated:
        return redirect('home')

    supplier = get_object_or_404(
        Supplier,
        id=supplier_id,
        user=request.user
    )

    if request.method == 'POST':
        supplier.delete()

        messages.success(
            request,
            'Supplier deleted successfully.'
        )

        return redirect('suppliers')

    return render(request, 'delete_supplier.html', {
        'supplier': supplier
    })


def reports(request):
    if not request.user.is_authenticated:
        return redirect('home')

    return render(request, 'reports.html')



def sales_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    sales = Sale.objects.filter(
        user=request.user
    ).select_related(
        'product',
        'customer'
    ).order_by('-sale_date')

    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    if start_date:
        sales = sales.filter(
            sale_date__date__gte=start_date
        )

    if end_date:
        sales = sales.filter(
            sale_date__date__lte=end_date
        )

    total_sales = sales.count()

    total_revenue = sales.aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    return render(request, 'sales_report.html', {
        'sales': sales,
        'total_sales': total_sales,
        'total_revenue': total_revenue,
        'start_date': start_date,
        'end_date': end_date,
    })


def inventory_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    products = Product.objects.filter(
        user=request.user
    ).select_related(
        'supplier'
    ).order_by('name')

    total_products = products.count()

    total_stock = products.aggregate(
        total=Sum('quantity')
    )['total'] or 0

    total_inventory_value = 0

    for product in products:
        product.inventory_value = product.quantity * product.selling_price
        total_inventory_value += product.inventory_value

    low_stock_products = products.filter(
        quantity__lte=5
    )

    return render(request, 'inventory_report.html', {
        'products': products,
        'total_products': total_products,
        'total_stock': total_stock,
        'total_inventory_value': total_inventory_value,
        'low_stock_products': low_stock_products,
    })


def customer_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    customers = Customer.objects.filter(
        user=request.user
    ).annotate(
        total_orders=Count('sale'),
        total_purchases=Sum('sale__total_amount'),
        last_purchase=Max('sale__sale_date')
    ).order_by('name')

    total_customers = customers.count()

    total_orders = Sale.objects.filter(
        user=request.user,
        customer__isnull=False
    ).count()

    total_purchases = Sale.objects.filter(
        user=request.user,
        customer__isnull=False
    ).aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    return render(request, 'customer_report.html', {
        'customers': customers,
        'total_customers': total_customers,
        'total_orders': total_orders,
        'total_purchases': total_purchases,
    })


def supplier_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    suppliers = Supplier.objects.filter(
        user=request.user
    ).annotate(
        total_products=Count('product'),
        total_stock=Sum('product__quantity')
    ).order_by('name')

    total_suppliers = suppliers.count()

    total_products = Product.objects.filter(
        user=request.user,
        supplier__isnull=False
    ).count()

    total_stock = Product.objects.filter(
        user=request.user,
        supplier__isnull=False
    ).aggregate(
        total=Sum('quantity')
    )['total'] or 0

    return render(request, 'supplier_report.html', {
        'suppliers': suppliers,
        'total_suppliers': total_suppliers,
        'total_products': total_products,
        'total_stock': total_stock,
    })


def revenue_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    sales = Sale.objects.filter(
        user=request.user
    ).select_related(
        'product',
        'customer'
    ).order_by('-sale_date')

    # Summary
    total_sales = sales.count()

    total_revenue = sales.aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    average_sale = sales.aggregate(
        average=Avg('total_amount')
    )['average'] or 0

    # Payment status summary
    paid_sales = sales.filter(
        payment_status='Paid'
    ).count()

    pending_sales = sales.filter(
        payment_status='Pending'
    ).count()

    partial_sales = sales.filter(
        payment_status='Partial'
    ).count()

    # Daily revenue
    daily_revenue = {}

    for sale in sales:
        date = sale.sale_date.date()

        if date not in daily_revenue:
            daily_revenue[date] = 0

        daily_revenue[date] += sale.total_amount

    daily_revenue = sorted(
        daily_revenue.items(),
        reverse=True
    )

    return render(request, 'revenue_report.html', {
        'sales': sales,
        'total_sales': total_sales,
        'total_revenue': total_revenue,
        'average_sale': average_sale,
        'paid_sales': paid_sales,
        'pending_sales': pending_sales,
        'partial_sales': partial_sales,
        'daily_revenue': daily_revenue,
    })

def pending_payments_report(request):
    if not request.user.is_authenticated:
        return redirect('home')

    pending_sales = Sale.objects.filter(
        user=request.user,
        payment_status__in=['Pending', 'Partial']
    ).select_related(
        'product',
        'customer'
    ).order_by('-sale_date')

    total_pending = pending_sales.aggregate(
        total=Sum('total_amount')
    )['total'] or 0

    pending_count = pending_sales.count()

    return render(request, 'pending_payments_report.html', {
        'pending_sales': pending_sales,
        'total_pending': total_pending,
        'pending_count': pending_count,
    })


def settings_view(request):
    if not request.user.is_authenticated:
        return redirect('home')

    return render(request, 'settings.html')


def update_profile(request):
    if not request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':

        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        email = request.POST.get('email', '').strip()

        user = request.user

        user.first_name = first_name
        user.last_name = last_name
        user.email = email

        user.save()

        messages.success(
            request,
            'Profile updated successfully.'
        )

    return redirect('settings')


def change_password(request):
    if not request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':

        form = PasswordChangeForm(
            request.user,
            request.POST
        )

        if form.is_valid():

            user = form.save()

            update_session_auth_hash(
                request,
                user
            )

            messages.success(
                request,
                'Password changed successfully.'
            )

            return redirect('settings')

        else:

            return render(
                request,
                'settings.html',
                {
                    'password_form': form
                }
            )

    return redirect('settings')


def stock_history(request):
    if not request.user.is_authenticated:
        return redirect('home')

    movements = StockMovement.objects.filter(
        user=request.user
    ).select_related('product').order_by('-created_at')

    return render(request, 'stock_history.html', {
        'movements': movements
    })


def export_sales_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    sales_data = Sale.objects.filter(
        user=request.user
    ).select_related('product', 'customer').order_by('-sale_date')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sales Report"

    headers = [
        "Invoice",
        "Product",
        "Customer",
        "Quantity",
        "Selling Price",
        "Total Amount",
        "Payment Status",
        "Date"
    ]

    worksheet.append(headers)

    for sale in sales_data:
        worksheet.append([
            f"INV-{sale.id}",
            sale.product.name,
            sale.customer.name if sale.customer else sale.customer_name,
            sale.quantity,
            float(sale.selling_price),
            float(sale.total_amount),
            sale.payment_status,
            sale.sale_date.strftime("%d-%m-%Y %H:%M")
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = 'attachment; filename="sales_report.xlsx"'

    workbook.save(response)

    return response


def export_inventory_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    products = Product.objects.filter(
        user=request.user
    ).select_related('supplier').order_by('name')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Inventory Report"

    headers = [
        "Product",
        "SKU",
        "Category",
        "Supplier",
        "Quantity",
        "Purchase Price",
        "Selling Price",
        "Stock Status"
    ]

    worksheet.append(headers)

    for product in products:

        if product.quantity <= 5:
            stock_status = "Low Stock"
        else:
            stock_status = "Available"

        worksheet.append([
            product.name,
            product.sku,
            product.category,
            product.supplier.name if product.supplier else "",
            product.quantity,
            float(product.purchase_price),
            float(product.selling_price),
            stock_status
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="inventory_report.xlsx"'
    )

    workbook.save(response)

    return response


def export_customer_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    customers_data = Customer.objects.filter(
        user=request.user
    ).order_by('name')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Customer Report"

    headers = [
        "Customer",
        "Email",
        "Phone",
        "Total Purchases",
        "Created Date"
    ]

    worksheet.append(headers)

    for customer in customers_data:
        total_purchases = customer.sale_set.aggregate(
            total=Sum('total_amount')
        )['total'] or 0

        worksheet.append([
            customer.name,
            customer.email,
            customer.phone,
            float(total_purchases),
            customer.created_at.strftime("%d-%m-%Y")
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="customer_report.xlsx"'
    )

    workbook.save(response)

    return response


def export_supplier_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    suppliers_data = Supplier.objects.filter(
        user=request.user
    ).order_by('name')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Supplier Report"

    headers = [
        "Supplier",
        "Company",
        "Email",
        "Phone",
        "Address",
        "Created Date"
    ]

    worksheet.append(headers)

    for supplier in suppliers_data:
        worksheet.append([
            supplier.name,
            supplier.company,
            supplier.email,
            supplier.phone,
            supplier.address,
            supplier.created_at.strftime("%d-%m-%Y")
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="supplier_report.xlsx"'
    )

    workbook.save(response)

    return response


def export_revenue_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    sales_data = Sale.objects.filter(
        user=request.user
    ).order_by('-sale_date')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Revenue Report"

    headers = [
        "Invoice",
        "Product",
        "Customer",
        "Quantity",
        "Selling Price",
        "Revenue",
        "Payment Status",
        "Sale Date"
    ]

    worksheet.append(headers)

    for sale in sales_data:
        worksheet.append([
            f"INV-{sale.id}",
            sale.product.name,
            sale.customer.name if sale.customer else sale.customer_name,
            sale.quantity,
            float(sale.selling_price),
            float(sale.total_amount),
            sale.payment_status,
            sale.sale_date.strftime("%d-%m-%Y")
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="revenue_report.xlsx"'
    )

    workbook.save(response)

    return response


def export_pending_payments_excel(request):
    if not request.user.is_authenticated:
        return redirect('home')

    pending_sales = Sale.objects.filter(
        user=request.user,
        payment_status__in=['Pending', 'Partial']
    ).select_related(
        'product',
        'customer'
    ).order_by('-sale_date')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Pending Payments"

    headers = [
        "Invoice",
        "Customer",
        "Product",
        "Quantity",
        "Total Amount",
        "Payment Status",
        "Sale Date"
    ]

    worksheet.append(headers)

    for sale in pending_sales:
        worksheet.append([
            f"INV-{sale.id}",
            sale.customer.name if sale.customer else sale.customer_name,
            sale.product.name,
            sale.quantity,
            float(sale.total_amount),
            sale.payment_status,
            sale.sale_date.strftime("%d-%m-%Y")
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="pending_payments.xlsx"'
    )

    workbook.save(response)

    return response